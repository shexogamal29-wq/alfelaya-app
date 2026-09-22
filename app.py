import asyncio
import io
import re
from datetime import datetime, timedelta
from typing import List, Optional
from fastapi import FastAPI, HTTPException, UploadFile, File, WebSocket, WebSocketDisconnect
from fastapi.responses import FileResponse, Response
from pydantic import BaseModel
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from sqlalchemy import Boolean, Column, DateTime, Float, Integer, String, create_engine, desc
from sqlalchemy.orm import declarative_base, sessionmaker

# --- 1. إعداد قاعدة البيانات ---
DATABASE_URL = "sqlite:///./real_estate.db"
engine = create_engine(DATABASE_URL, connect_args={"check_same_thread": False})
SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine, expire_on_commit=False)
Base = declarative_base()

# --- 2. الجداول ---
class SystemSetting(Base):
    __tablename__ = "system_settings"
    id = Column(Integer, primary_key=True, index=True)
    key = Column(String, unique=True, index=True)
    value = Column(String)

class User(Base):
    __tablename__ = "users"
    id = Column(Integer, primary_key=True, index=True)
    username = Column(String, unique=True, index=True)
    full_name = Column(String)
    password = Column(String)
    role = Column(String, default="sales")  # manager, supervisor, sales
    created_at = Column(DateTime, default=datetime.utcnow)

class Unit(Base):
    __tablename__ = "units"
    id = Column(Integer, primary_key=True, index=True)
    unit_code = Column(String, unique=True, index=True)
    building = Column(String, default="1")
    category = Column(String, default="سكني")        # سكني، تجاري، إداري، طبي
    floor = Column(Integer)
    unit_number = Column(Integer, default=1)
    area_sqm = Column(Float, default=120.0)
    meter_price = Column(Float, default=0.0)
    price = Column(Float)
    discount = Column(Float, default=0.0)
    final_price = Column(Float, nullable=True)
    status = Column(String, default="available")    # available, hold, sold
    locked_by = Column(String, nullable=True)
    locked_by_user_id = Column(Integer, nullable=True)
    hold_expires_at = Column(DateTime, nullable=True)
    is_blocked = Column(Boolean, default=False)
    hidden_from = Column(String, default="none")    # none, sales, sales_and_supervisor

class AuditLog(Base):
    __tablename__ = "audit_logs"
    id = Column(Integer, primary_key=True, index=True)
    timestamp = Column(DateTime, default=datetime.utcnow)
    user_name = Column(String)
    user_role = Column(String)
    action = Column(String)
    unit_code = Column(String, nullable=True)
    details = Column(String)

class AllowedDevice(Base):
    __tablename__ = "allowed_devices"
    id = Column(Integer, primary_key=True, index=True)
    device_token = Column(String, unique=True, index=True)
    device_name = Column(String)
    user_name = Column(String)
    is_approved = Column(Boolean, default=True)

Base.metadata.create_all(bind=engine)

def generate_smart_code(building: str, floor: int, unit_no: int) -> str:
    digits = re.findall(r'\d+', str(building))
    b_prefix = digits[0] if digits else "1"
    return f"{b_prefix}{floor}{unit_no:02d}"

def log_action(db, user_name: str, user_role: str, action: str, details: str, unit_code: Optional[str] = None):
    new_log = AuditLog(
        user_name=user_name,
        user_role=user_role,
        action=action,
        unit_code=unit_code,
        details=details
    )
    db.add(new_log)
    db.commit()

def parse_arabic_number(val):
    if val is None:
        return 0.0
    s = str(val).strip()
    eastern_to_western = str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789")
    s = s.translate(eastern_to_western)
    s = s.replace(",", "").replace("،", "").replace(" ", "").replace("\xa0", "")
    try:
        return float(s)
    except:
        return 0.0

# --- 3. WebSocket Manager ---
class ConnectionManager:
    def __init__(self):
        self.active_connections: List[WebSocket] = []

    async def connect(self, websocket: WebSocket):
        await websocket.accept()
        self.active_connections.append(websocket)

    def disconnect(self, websocket: WebSocket):
        if websocket in self.active_connections:
            self.active_connections.remove(websocket)

    async def broadcast(self, message: dict):
        for connection in self.active_connections:
            try:
                await connection.send_json(message)
            except:
                pass

manager = ConnectionManager()
app = FastAPI(title="Real Estate Enterprise System")

# --- 4. Schemas ---
class LoginRequest(BaseModel):
    username: str
    password: str
    device_token: str
    device_name: str

class CreateUserRequest(BaseModel):
    manager_user_id: int
    username: str
    password: str
    full_name: str
    role: str

class HoldUnitRequest(BaseModel):
    unit_id: int
    user_id: int

class ExtendHoldRequest(BaseModel):
    unit_id: int
    user_id: int
    additional_hours: int

class UnitActionRequest(BaseModel):
    unit_id: int
    user_id: int

class MarkSoldRequest(BaseModel):
    unit_id: int
    user_id: int
    discount_amount: float = 0.0

class BlockUnitsRequest(BaseModel):
    manager_user_id: int
    scope: str
    target_building: str
    target_floor: Optional[int] = None
    target_unit_id: Optional[int] = None
    is_blocked: bool
    hidden_from: str

class CreateUnitRequest(BaseModel):
    user_id: int
    building: str
    category: str
    floor: int
    unit_number: int
    area_sqm: float
    meter_price: float

class UpdateHoldDurationRequest(BaseModel):
    manager_user_id: int
    hold_hours: int

# --- 5. فاحص انتهاء الحجوزات التلقائي ---
async def auto_release_expired_holds():
    while True:
        try:
            db = SessionLocal()
            now = datetime.utcnow()
            expired_units = db.query(Unit).filter(
                Unit.status == "hold",
                Unit.hold_expires_at != None,
                Unit.hold_expires_at <= now
            ).all()

            for unit in expired_units:
                old_holder = unit.locked_by
                u_code = unit.unit_code
                unit.status = "available"
                unit.locked_by = None
                unit.locked_by_user_id = None
                unit.hold_expires_at = None
                db.commit()

                log_action(db, "النظام التلقائي", "System", "فك حجز تلقائي", f"انتهت مهلة الحجز المحددة للشقة {u_code} (كانت مع {old_holder}) وتم إرجاعها للمتاح", u_code)
                await manager.broadcast({"event": "reload_needed"})
            db.close()
        except Exception as e:
            print(f"Auto-release error: {e}")
        await asyncio.sleep(30)

@app.on_event("startup")
def startup_event():
    db = SessionLocal()
    try:
        # إعداد مدة الحجز الافتراضية: 48 ساعة
        setting = db.query(SystemSetting).filter(SystemSetting.key == "default_hold_hours").first()
        if not setting:
            db.add(SystemSetting(key="default_hold_hours", value="48"))
            db.commit()

        # المستخدمين الافتراضيين
        if db.query(User).count() == 0:
            db.add_all([
                User(username="admin", full_name="مدير المبيعات", password="123", role="manager"),
                User(username="super1", full_name="كريم مشرف", password="123", role="supervisor"),
                User(username="sales1", full_name="أحمد مبيعات", password="123", role="sales"),
                User(username="sales2", full_name="محمود مبيعات", password="123", role="sales"),
            ])
            db.commit()
    finally:
        db.close()
    asyncio.create_task(auto_release_expired_holds())

@app.get("/")
def home():
    return FileResponse("index.html")

# --- 6. تسجيل الدخول ---
@app.post("/auth/login")
def login(req: LoginRequest):
    db = SessionLocal()
    try:
        user = db.query(User).filter(User.username == req.username, User.password == req.password).first()
        if not user:
            raise HTTPException(status_code=401, detail="اسم المستخدم أو كلمة المرور غير صحيحة")

        user_info = {
            "id": int(user.id),
            "username": str(user.username),
            "full_name": str(user.full_name),
            "role": str(user.role)
        }

        device = db.query(AllowedDevice).filter(AllowedDevice.device_token == req.device_token).first()
        if not device:
            new_device = AllowedDevice(
                device_token=req.device_token,
                device_name=req.device_name,
                user_name=user_info["full_name"],
                is_approved=True
            )
            db.add(new_device)
            db.commit()

        log_action(db, user.full_name, user.role, "تسجيل دخول", f"دخول من جهاز: {req.device_name}")
        return {"status": "success", "user": user_info}
    finally:
        db.close()

# --- 7. إعدادات مدة الحجز ---
@app.get("/settings/hold-duration")
def get_hold_duration():
    db = SessionLocal()
    try:
        setting = db.query(SystemSetting).filter(SystemSetting.key == "default_hold_hours").first()
        hours = int(setting.value) if setting else 48
        return {"hold_hours": hours}
    finally:
        db.close()

@app.post("/settings/hold-duration")
def update_hold_duration(req: UpdateHoldDurationRequest):
    db = SessionLocal()
    try:
        mgr = db.query(User).filter(User.id == req.manager_user_id, User.role == "manager").first()
        if not mgr:
            raise HTTPException(status_code=403, detail="تعديل مدة الحجز صلاحية للمدير فقط")

        setting = db.query(SystemSetting).filter(SystemSetting.key == "default_hold_hours").first()
        if not setting:
            setting = SystemSetting(key="default_hold_hours", value=str(req.hold_hours))
            db.add(setting)
        else:
            setting.value = str(req.hold_hours)
        db.commit()

        log_action(db, mgr.full_name, mgr.role, "تعديل إعدادات النظام", f"تعديل مدة الحجز الافتراضية لتصبح {req.hold_hours} ساعة")
        return {"status": "success", "hold_hours": req.hold_hours}
    finally:
        db.close()

# --- 8. إدارة الوحدات مع الحجب الذكي ---
@app.get("/units")
def get_units(user_id: Optional[int] = None):
    db = SessionLocal()
    try:
        role = "guest"
        if user_id:
            user = db.query(User).filter(User.id == user_id).first()
            if user:
                role = user.role

        all_units = db.query(Unit).all()
        filtered = []

        for u in all_units:
            if u.is_blocked:
                if role == "sales" and u.hidden_from in ["sales", "sales_and_supervisor"]:
                    continue
                if role == "supervisor" and u.hidden_from == "sales_and_supervisor":
                    continue

            filtered.append({
                "id": u.id,
                "unit_code": u.unit_code,
                "building": u.building,
                "category": u.category,
                "floor": u.floor,
                "unit_number": u.unit_number,
                "area_sqm": u.area_sqm,
                "meter_price": u.meter_price,
                "price": u.price,
                "discount": u.discount,
                "final_price": u.final_price,
                "status": u.status,
                "locked_by": u.locked_by,
                "hold_expires_at": u.hold_expires_at.isoformat() if u.hold_expires_at else None,
                "is_blocked": u.is_blocked,
                "hidden_from": u.hidden_from
            })
        return filtered
    finally:
        db.close()

@app.post("/units/add")
async def add_unit(req: CreateUnitRequest):
    db = SessionLocal()
    try:
        user = db.query(User).filter(User.id == req.user_id).first()
        if not user or user.role not in ["manager", "supervisor"]:
            raise HTTPException(status_code=403, detail="الصلاحية للمدير والمشرف فقط")

        code = generate_smart_code(req.building, req.floor, req.unit_number)
        existing = db.query(Unit).filter(Unit.unit_code == code).first()
        if existing:
            raise HTTPException(status_code=400, detail=f"كود الوحدة {code} مسجل مسبقاً")

        total_price = req.area_sqm * req.meter_price
        new_u = Unit(
            unit_code=code,
            building=req.building,
            category=req.category,
            floor=req.floor,
            unit_number=req.unit_number,
            area_sqm=req.area_sqm,
            meter_price=req.meter_price,
            price=total_price,
            status="available"
        )
        db.add(new_u)
        db.commit()

        log_action(db, user.full_name, user.role, "إضافة وحدة", f"إضافة الوحدة {code} ({req.category}) بـ {req.building}", code)
        await manager.broadcast({"event": "reload_needed"})
        return {"status": "success", "unit_code": code}
    finally:
        db.close()

@app.post("/units/block-visibility")
async def set_block_visibility(req: BlockUnitsRequest):
    db = SessionLocal()
    try:
        mgr = db.query(User).filter(User.id == req.manager_user_id, User.role == "manager").first()
        if not mgr:
            raise HTTPException(status_code=403, detail="الصلاحية للمدير فقط")

        q = db.query(Unit)
        target_name = ""

        if req.scope == "unit" and req.target_unit_id:
            q = q.filter(Unit.id == req.target_unit_id)
            target_name = f"الوحدة رقم {req.target_unit_id}"
        elif req.scope == "floor" and req.target_floor is not None:
            q = q.filter(Unit.building == req.target_building, Unit.floor == req.target_floor)
            target_name = f"الدور {req.target_floor} بـ {req.target_building}"
        elif req.scope == "building":
            q = q.filter(Unit.building == req.target_building)
            target_name = f"البرج {req.target_building} بالكامل"

        units = q.all()
        for u in units:
            u.is_blocked = req.is_blocked
            u.hidden_from = req.hidden_from if req.is_blocked else "none"
        db.commit()

        action_txt = "غلق وحجب عن " + ("السيلز والمشرف" if req.hidden_from == "sales_and_supervisor" else "السيلز فقط") if req.is_blocked else "فتح وإتاحة الطرح"
        log_action(db, mgr.full_name, mgr.role, "تعديل حجب الطرح", f"تم {action_txt} لـ {target_name}")

        await manager.broadcast({"event": "reload_needed"})
        return {"status": "success", "count": len(units)}
    finally:
        db.close()

@app.post("/units/hold")
async def hold_unit(req: HoldUnitRequest):
    db = SessionLocal()
    try:
        user = db.query(User).filter(User.id == req.user_id).first()
        if not user:
            raise HTTPException(status_code=404, detail="المستخدم غير موجود")

        unit = db.query(Unit).filter(Unit.id == req.unit_id).first()
        if not unit or unit.status != "available":
            raise HTTPException(status_code=400, detail="الوحدة غير متاحة للحجز")

        setting = db.query(SystemSetting).filter(SystemSetting.key == "default_hold_hours").first()
        hours = int(setting.value) if setting else 48

        expires_at = datetime.utcnow() + timedelta(hours=hours)
        unit.status = "hold"
        unit.locked_by = user.full_name
        unit.locked_by_user_id = user.id
        unit.hold_expires_at = expires_at
        db.commit()

        log_action(db, user.full_name, user.role, "حجز مؤقت", f"حجز لمدة {hours} ساعة تلقائياً ينتهي في {expires_at.strftime('%Y-%m-%d %H:%M')}", unit.unit_code)
        await manager.broadcast({"event": "reload_needed"})
        return {"status": "success"}
    finally:
        db.close()

@app.post("/units/extend-hold")
async def extend_hold(req: ExtendHoldRequest):
    db = SessionLocal()
    try:
        user = db.query(User).filter(User.id == req.user_id).first()
        if not user or user.role not in ["manager", "supervisor"]:
            raise HTTPException(status_code=403, detail="تمديد الحجز للمشرف والمدير فقط")

        unit = db.query(Unit).filter(Unit.id == req.unit_id).first()
        if not unit or unit.status != "hold":
            raise HTTPException(status_code=400, detail="الوحدة ليست قيد الحجز")

        base_time = unit.hold_expires_at if (unit.hold_expires_at and unit.hold_expires_at > datetime.utcnow()) else datetime.utcnow()
        unit.hold_expires_at = base_time + timedelta(hours=req.additional_hours)
        db.commit()

        log_action(db, user.full_name, user.role, "تمديد حجز", f"تمديد بمقدار {req.additional_hours} ساعة إضافية", unit.unit_code)
        await manager.broadcast({"event": "reload_needed"})
        return {"status": "success"}
    finally:
        db.close()

@app.post("/units/release")
async def release_unit(req: UnitActionRequest):
    db = SessionLocal()
    try:
        user = db.query(User).filter(User.id == req.user_id).first()
        if not user or user.role not in ["manager", "supervisor"]:
            raise HTTPException(status_code=403, detail="إلغاء الحجز للمشرف والمدير فقط")

        unit = db.query(Unit).filter(Unit.id == req.unit_id).first()
        if not unit:
            raise HTTPException(status_code=404, detail="الوحدة غير موجودة")

        old_holder = unit.locked_by
        unit.status = "available"
        unit.locked_by = None
        unit.locked_by_user_id = None
        unit.hold_expires_at = None
        db.commit()

        log_action(db, user.full_name, user.role, "إلغاء حجز", f"إلغاء حجز الشقة (كانت مع {old_holder})", unit.unit_code)
        await manager.broadcast({"event": "reload_needed"})
        return {"status": "success"}
    finally:
        db.close()

@app.post("/units/mark-sold")
async def mark_sold(req: MarkSoldRequest):
    db = SessionLocal()
    try:
        user = db.query(User).filter(User.id == req.user_id, User.role == "manager").first()
        if not user:
            raise HTTPException(status_code=403, detail="اعتماد البيع وتطبيق الخصم حصري للمدير فقط")

        unit = db.query(Unit).filter(Unit.id == req.unit_id).first()
        if not unit:
            raise HTTPException(status_code=404, detail="الوحدة غير موجودة")

        unit.status = "sold"
        unit.discount = req.discount_amount
        unit.final_price = max(0.0, unit.price - req.discount_amount)
        unit.hold_expires_at = None
        db.commit()

        log_action(
            db, user.full_name, user.role, "اعتماد بيع نهائي",
            f"تم البيع بسعر نهائي {unit.final_price:,.0f} ج.م (خصم: {unit.discount:,.0f} ج.م) بواسطة: {unit.locked_by or 'مباشر'}",
            unit.unit_code
        )

        await manager.broadcast({"event": "reload_needed"})
        return {"status": "success", "final_price": unit.final_price}
    finally:
        db.close()

# --- 9. استيراد الإكسيل الذكي المرن ---
@app.post("/units/import-excel")
async def import_excel(manager_user_id: int, file: UploadFile = File(...)):
    db = SessionLocal()
    try:
        user = db.query(User).filter(User.id == manager_user_id, User.role == "manager").first()
        if not user:
            raise HTTPException(status_code=403, detail="استيراد البيانات مسموح للمدير فقط")

        contents = await file.read()
        wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
        ws = wb.active

        header_row = [str(c).strip() if c is not None else "" for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True), [])]
        col_map = {}
        for idx, name in enumerate(header_row):
            clean = name.replace(" ", "").replace("أ", "ا").replace("إ", "ا")
            if "برج" in clean or "عماره" in clean: col_map["b"] = idx
            elif "دور" in clean: col_map["flr"] = idx
            elif "شقه" in clean or "وحده" in clean: col_map["unit"] = idx
            elif "تصنيف" in clean or "نوع" in clean: col_map["cat"] = idx
            elif "مساح" in clean: col_map["area"] = idx
            elif "متر" in clean or "سعر" in clean: col_map["price"] = idx

        b_idx = col_map.get("b", 0)
        flr_idx = col_map.get("flr", 1)
        unit_idx = col_map.get("unit", 2)
        cat_idx = col_map.get("cat", 3)
        area_idx = col_map.get("area", 4)
        price_idx = col_map.get("price", 5)

        imported_count = 0
        skipped_count = 0

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or all(v is None for v in row):
                continue
            try:
                b_raw = row[b_idx] if b_idx < len(row) else 1
                flr_raw = row[flr_idx] if flr_idx < len(row) else None
                unit_raw = row[unit_idx] if unit_idx < len(row) else None
                cat_raw = row[cat_idx] if cat_idx < len(row) and row[cat_idx] else "سكني"
                area_raw = row[area_idx] if area_idx < len(row) else None
                price_raw = row[price_idx] if price_idx < len(row) else None

                if flr_raw is None or unit_raw is None or area_raw is None or price_raw is None:
                    skipped_count += 1
                    continue

                b_name = str(b_raw).strip()
                flr = int(parse_arabic_number(flr_raw))
                u_num = int(parse_arabic_number(unit_raw))
                cat = str(cat_raw).strip()
                area = parse_arabic_number(area_raw)
                m_price = parse_arabic_number(price_raw)

                if area <= 0 or m_price <= 0:
                    skipped_count += 1
                    continue

                total_price = area * m_price
                code = generate_smart_code(b_name, flr, u_num)

                existing = db.query(Unit).filter(Unit.unit_code == code).first()
                if existing:
                    existing.building = b_name
                    existing.floor = flr
                    existing.unit_number = u_num
                    existing.category = cat
                    existing.area_sqm = area
                    existing.meter_price = m_price
                    existing.price = total_price
                else:
                    new_u = Unit(
                        unit_code=code,
                        building=b_name,
                        category=cat,
                        floor=flr,
                        unit_number=u_num,
                        area_sqm=area,
                        meter_price=m_price,
                        price=total_price,
                        status="available"
                    )
                    db.add(new_u)
                imported_count += 1
            except:
                skipped_count += 1

        db.commit()
        log_action(db, user.full_name, user.role, "استيراد إكسيل", f"تم استيراد {imported_count} وحدة بنجاح")
        await manager.broadcast({"event": "reload_needed"})
        return {"status": "success", "imported": imported_count, "skipped": skipped_count}
    finally:
        db.close()

# --- 10. تصدير التقارير (محمي ضد تعليق صفحات المتصفح) ---
@app.get("/reports/export-excel")
def export_excel_report(report_type: str, user_id: int):
    db = SessionLocal()
    try:
        user = db.query(User).filter(User.id == user_id).first()
        if not user or user.role not in ["manager", "supervisor"]:
            raise HTTPException(status_code=403, detail="الصلاحية إدارية")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.views.sheetView[0].rightToLeft = True

        header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        center_align = Alignment(horizontal="center", vertical="center")

        if report_type == "sold":
            ws.title = "الوحدات المباعة"
            headers = ["كود الوحدة", "العمارة/البرج", "التصنيف", "الدور", "رقم الشقة", "المساحة (م²)", "سعر المتر", "السعر الأساسي", "الخصم", "سعر البيع النهائي", "السيلز الحاجز"]
            units = db.query(Unit).filter(Unit.status == "sold").all()
        elif report_type == "hold":
            ws.title = "الوحدات المحجوزة حالياً"
            headers = ["كود الوحدة", "العمارة/البرج", "التصنيف", "الدور", "رقم الشقة", "المساحة (م²)", "سعر المتر", "إجمالي السعر", "اسم السيلز", "موعد الانتهاء"]
            units = db.query(Unit).filter(Unit.status == "hold").all()
        else:
            ws.title = "الوحدات المتاحة"
            headers = ["كود الوحدة", "العمارة/البرج", "التصنيف", "الدور", "رقم الشقة", "المساحة (م²)", "سعر المتر", "السعر المطلوب"]
            units = db.query(Unit).filter(Unit.status == "available").all()

        ws.append(headers)
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align

        for u in units:
            if report_type == "sold":
                ws.append([
                    u.unit_code, u.building, u.category, u.floor, u.unit_number, u.area_sqm,
                    f"{u.meter_price:,.0f} ج.م", f"{u.price:,.0f} ج.م", f"{u.discount:,.0f} ج.م",
                    f"{(u.final_price or u.price):,.0f} ج.م", u.locked_by or "مباشر"
                ])
            elif report_type == "hold":
                exp = u.hold_expires_at.strftime("%Y-%m-%d %H:%M") if u.hold_expires_at else "-"
                ws.append([
                    u.unit_code, u.building, u.category, u.floor, u.unit_number, u.area_sqm,
                    f"{u.meter_price:,.0f} ج.م", f"{u.price:,.0f} ج.م", u.locked_by or "-", exp
                ])
            else:
                ws.append([
                    u.unit_code, u.building, u.category, u.floor, u.unit_number,
                    u.area_sqm, f"{u.meter_price:,.0f} ج.م", f"{u.price:,.0f} ج.م"
                ])

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f"report_{report_type}_{datetime.utcnow().strftime('%Y%m%d')}.xlsx"
        return Response(
            content=output.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f"attachment; filename={filename}",
                "Cache-Control": "no-cache, no-store, must-revalidate",
                "Pragma": "no-cache",
                "Expires": "0"
            }
        )
    finally:
        db.close()

# --- 11. إدارة الموظفين وسجل العمليات ---
@app.get("/users/list")
def list_users(user_id: int):
    db = SessionLocal()
    try:
        u = db.query(User).filter(User.id == user_id, User.role == "manager").first()
        if not u: raise HTTPException(status_code=403, detail="الصلاحية للمدير فقط")
        users = db.query(User).order_by(User.id).all()
        return [{"id": x.id, "username": x.username, "full_name": x.full_name, "role": x.role} for x in users]
    finally:
        db.close()

@app.post("/users/add")
def create_user(req: CreateUserRequest):
    db = SessionLocal()
    try:
        mgr = db.query(User).filter(User.id == req.manager_user_id, User.role == "manager").first()
        if not mgr: raise HTTPException(status_code=403, detail="الصلاحية للمدير فقط")
        if db.query(User).filter(User.username == req.username).first():
            raise HTTPException(status_code=400, detail="اسم المستخدم مسجل مسبقاً")

        new_u = User(username=req.username, password=req.password, full_name=req.full_name, role=req.role)
        db.add(new_u)
        db.commit()
        log_action(db, mgr.full_name, mgr.role, "إضافة موظف", f"تم إنشاء حساب للموظف {new_u.full_name} برتبة {new_u.role}")
        return {"status": "success", "message": f"تمت إضافة {new_u.full_name} بنجاح"}
    finally:
        db.close()

@app.delete("/users/{user_id_to_delete}")
def delete_user(user_id_to_delete: int, manager_id: int):
    db = SessionLocal()
    try:
        mgr = db.query(User).filter(User.id == manager_id, User.role == "manager").first()
        if not mgr: raise HTTPException(status_code=403, detail="الصلاحية للمدير فقط")
        if user_id_to_delete == manager_id:
            raise HTTPException(status_code=400, detail="لا يمكن حذف الحساب الحالي للمدير")
        target = db.query(User).filter(User.id == user_id_to_delete).first()
        if target:
            db.delete(target)
            db.commit()
        return {"status": "success"}
    finally:
        db.close()

@app.get("/admin/logs")
def get_audit_logs(user_id: int):
    db = SessionLocal()
    try:
        mgr = db.query(User).filter(User.id == user_id, User.role == "manager").first()
        if not mgr: raise HTTPException(status_code=403, detail="الصلاحية للمدير فقط")
        logs = db.query(AuditLog).order_by(desc(AuditLog.id)).limit(100).all()
        return [{"time": l.timestamp.strftime("%Y-%m-%d %H:%M:%S"), "user": l.user_name, "role": l.user_role, "action": l.action, "unit": l.unit_code or "-", "details": l.details} for l in logs]
    finally:
        db.close()

# --- 12. WebSocket ---
@app.websocket("/ws")
async def websocket_endpoint(websocket: WebSocket):
    await manager.connect(websocket)
    try:
        while True:
            await websocket.receive_text()
    except WebSocketDisconnect:
        manager.disconnect(websocket)